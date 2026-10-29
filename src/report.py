"""
Step 7 - Generate the final project report (Word .docx) from the saved results.

All numbers and most interpretation sentences are computed from outputs/results/*,
so re-running the pipeline on the real PHEME data regenerates a correct report.
"""
from __future__ import annotations

import json
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .features import FEATURE_DESCRIPTIONS, FEATURE_GROUPS

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)
WARN = RGBColor(0xB0, 0x3A, 0x0E)

AUTHORS = [("Palivela Sai Hruthik", "AM.SC.U4AIE23141"), ("Nallapaneni Vaibhav", "AM.SC.U4AIE23135")]
COURSE = "23 CSE 356 Social Network Analysis"
TITLE = "Early Viral Misinformation Detection Using Social Network Analysis"

REFERENCES = [
    'P. N. Ahmad, A. M. Shah, K. Y. Lee, and W. Muhammad, "Misinformation detection on online social networks using pretrained language models," Information Processing & Management, vol. 63, no. 1, art. no. 104342, Jan. 2026.',
    'M. Sikosana, S. Maudsley-Barton, and O. Ajao, "Analysing health misinformation with advanced centrality metrics in online social networks," PLOS Digital Health, vol. 4, no. 6, art. no. e0000888, Jun. 2025.',
    'W. Li, S. J. Cranmer, Z. Zheng, and P. J. Mucha, "Infectivity enhances prediction of viral cascades in Twitter," PLOS ONE, vol. 14, no. 4, art. no. e0214453, Apr. 2019.',
    'L. Weng, F. Menczer, and Y.-Y. Ahn, "Virality prediction and community structure in social networks," Scientific Reports, vol. 3, art. no. 2522, 2013.',
    'E. Kochkina, M. Liakata, and A. Zubiaga, "All-in-one: Multi-task learning for rumour verification," in Proc. COLING, 2018. Dataset: PHEME dataset for Rumour Detection and Veracity Classification, figshare, doi:10.6084/m9.figshare.6392078.',
    'A. Zubiaga, M. Liakata, R. Procter, G. Wong Sak Hoi, and P. Tolmie, "Analysing how people orient to and spread rumours in social media by looking at conversational threads," PLOS ONE, vol. 11, no. 3, e0150989, 2016.',
    'S. Vosoughi, D. Roy, and S. Aral, "The spread of true and false news online," Science, vol. 359, no. 6380, pp. 1146-1151, 2018.',
    'S. Goel, A. Anderson, J. Hofman, and D. J. Watts, "The structural virality of online diffusion," Management Science, vol. 62, no. 1, pp. 180-196, 2016.',
    'V. D. Blondel, J.-L. Guillaume, R. Lambiotte, and E. Lefebvre, "Fast unfolding of communities in large networks," J. Stat. Mech., P10008, 2008.',
]


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------
def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _table(doc, header, rows, widths=None, font=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(font)
        _shade(c, "DCE8F7")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def _fig(doc, path, caption, width=15.5):
    if path and path.exists():
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        r = p.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _p(doc, text, bold_lead=None, size=11):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.size = Pt(size)
    r = p.add_run(text)
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6)
    return p


def _bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def _pct(x):
    return f"{x:.0%}"


# ---------------------------------------------------------------------------
def _key_numbers(metrics, text_base, imp, ttf, windows, mains):
    k = {}
    for task in metrics.task.unique():
        m = metrics[(metrics.task == task) & (metrics.model == mains[task])]
        a = m[m.feature_set == "all"].set_index("window")
        k[task] = {
            "auc": a["auc"].to_dict(), "f1": a["macro_f1"].to_dict(),
            "lo": a["auc_lo"].to_dict(), "hi": a["auc_hi"].to_dict(),
            "topo": m[m.feature_set == "topology"].set_index("window")["auc"].to_dict(),
            "meta": m[m.feature_set == "user_meta"].set_index("window")["auc"].to_dict(),
            "text_auc": float(text_base[text_base.task == task]["auc"].iloc[0]),
            "text_f1": float(text_base[text_base.task == task]["macro_f1"].iloc[0]),
            "best_model": {w: metrics[(metrics.task == task) & (metrics.feature_set == "all") &
                                      (metrics.window == w)].sort_values("auc").iloc[-1].model
                           for w in a.index},
            "top": {w: imp[(imp.task == task) & (imp.window == w)]
                    .sort_values("importance", ascending=False)["feature"].head(3).tolist()
                    for w in a.index},
        }
        full = k[task]["auc"]["full"]
        k[task]["retained"] = {w: (v - 0.5) / (full - 0.5) if full > 0.5 else float("nan")
                               for w, v in k[task]["auc"].items()}
        earliest = next((str(w) for w in windows if k[task]["auc"].get(str(w), 0) >= k[task]["text_auc"]),
                        None)
        k[task]["beats_text_from"] = earliest
    if ttf is not None and len(ttf):
        for key, yv in (("ttf", 1), ("ttf_fa", 0)):
            d = ttf[ttf["y"] == yv] if "y" in ttf else ttf
            vc = d["first_flag_window"].value_counts(normalize=True)
            cum, acc = {}, 0.0
            for w in [str(x) for x in windows]:
                acc += vc.get(w, 0.0)
                cum[w] = acc
            k[key] = cum
    return k


def build_report(paths, windows, out_file, is_sample: bool, main=None):
    R, F = paths["results"], paths["figures"]
    metrics = pd.read_csv(R / "metrics.csv", dtype={"window": str})
    text_base = pd.read_csv(R / "text_baseline.csv")
    imp = pd.read_csv(R / "importance.csv", dtype={"window": str})
    per_event = pd.read_csv(R / "per_event.csv", dtype={"window": str})
    summary = pd.read_csv(R / "dataset_summary.csv", index_col=0)
    ttf = pd.read_csv(R / "time_to_flag.csv") if (R / "time_to_flag.csv").exists() else None
    run_info = json.loads((R / "run_info.json").read_text())
    mains = main or run_info["main_model"]
    main = mains["A_false_vs_true"]
    nice = {"LogReg": "logistic regression", "RandomForest": "random forest",
            "GradBoost": "gradient boosting"}.get(main, main)
    K = _key_numbers(metrics, text_base, imp, ttf, windows, mains)
    A, B = K["A_false_vs_true"], K["B_rumour_vs_nonrumour"]
    W = [str(w) for w in windows]

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)

    # ---------------- title page ------------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COURSE + "\nProject Report")
    r.font.size = Pt(13)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name, roll in AUTHORS:
        p.add_run(f"{name}  ({roll})\n").font.size = Pt(12)
    p.add_run(f"\n{date.today():%d %B %Y}").font.size = Pt(11)

    if is_sample:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("SAMPLE RUN - numbers in this report come from SYNTHETIC data in the PHEME "
                      "format, used only to test the pipeline. Run `python run_all.py` on the real "
                      "PHEME dataset to regenerate this report with real results.")
        r.bold = True
        r.font.color.rgb = WARN
        r.font.size = Pt(10.5)

    doc.add_page_break()

    # ---------------- abstract --------------------------------------------------
    doc.add_heading("Abstract", level=1)
    _p(doc,
       "False rumours spread faster than corrections, yet most detection systems either read the "
       "text of a post or wait for the complete sharing cascade before deciding. We ask a different "
       "question: how early, and how accurately, can misinformation be flagged using only the "
       "structure of how it spreads? Using the PHEME rumour-veracity dataset "
       f"({int(summary.loc['TOTAL', 'total'])} Twitter conversation threads across "
       f"{len(summary) - 1} breaking-news events), we build a time-stamped reply cascade and a user "
       "interaction graph for every thread, truncate them at "
       f"{', '.join(w + ' min' for w in W[:-1])} and the full cascade, and extract "
       f"{sum(len(v) for v in FEATURE_GROUPS.values())} structure-only features (cascade shape, "
       "propagation speed, centrality, Louvain communities and non-text user metadata). Models are "
       "evaluated with leave-one-event-out cross-validation so every prediction is made on a news "
       f"event the model has never seen. For misinformation (false vs true rumours) the {nice} "
       f"model reaches AUC {A['auc']['10']:.2f} after only 10 minutes, {A['auc']['30']:.2f} after "
       f"30 minutes and {A['auc']['full']:.2f} on the full cascade, compared with "
       f"{A['text_auc']:.2f} for a text-only TF-IDF baseline. At 30 minutes the structural model "
       f"retains {_pct(A['retained']['30'])} of the full-cascade signal (above chance), showing that "
       "much of the information needed to flag a false rumour is available long before it goes viral.")

    # ---------------- 1. introduction -----------------------------------------
    doc.add_heading("1. Problem Statement and Motivation", level=1)
    _p(doc, "Fake news and rumours spread on social media faster than the truth [7]. By the time a "
            "claim is fact-checked, it has often reached thousands of people. Most detection systems "
            "look only at what is said - the words of the post - or at the complete, fully-grown "
            "sharing cascade. Both are problematic: wording is easy to fake, and waiting for the full "
            "cascade means the decision arrives after the damage is done.")
    _p(doc, "How a claim moves through the network - who reacts first, how fast, how deep the "
            "conversation goes and which communities it crosses - is visible within minutes, is harder "
            "to manipulate than text, and can act as an early-warning signal even when the content "
            "looks convincing. This project measures how much of that signal exists early.")

    doc.add_heading("2. Research Gap and Novelty", level=1)
    _bullets(doc, [
        ("Early + veracity combined. ", "Virality-prediction work [3, 4] shows structure is predictive "
         "within the first hour, but targets popularity, not truth. Veracity work mostly uses complete "
         "cascades. We apply early-window analysis to the veracity question."),
        ("Structure-only by design. ", "Many SNA misinformation systems fuse network features with "
         "large text/LLM models [1], hiding how much the network contributes. We exclude text and "
         "report a text baseline separately, isolating the network's contribution."),
        ("Time-resolved evaluation. ", "Instead of one accuracy on finished cascades, we report a "
         "performance-vs-time curve, directly quantifying the earliness/accuracy trade-off."),
        ("Realistic generalisation. ", "Leave-one-event-out evaluation tests the model on breaking "
         "news it has never seen, the situation an early-warning system actually faces."),
    ])

    doc.add_heading("3. Objectives", level=1)
    _bullets(doc, [
        "Build time-stamped network representations of each claim's spread for early windows "
        f"({', '.join(w + ' min' for w in W[:-1])}) and the full cascade.",
        "Identify which structure-only features (degree, clustering of the conversation, community "
        "bridging, propagation velocity, cascade depth/breadth) are most predictive of misinformation.",
        "Compare early-window structural detection with a full-cascade SNA baseline and a "
        "content-based (text) baseline, measuring how early and how accurately misinformation can be flagged.",
    ])

    # ---------------- 4. dataset ---------------------------------------------
    doc.add_heading("4. Dataset", level=1)
    _p(doc, "We use the PHEME dataset for Rumour Detection and Veracity Classification [5, 6]: Twitter "
            "conversation threads collected during nine breaking-news events. Each thread contains a "
            "source tweet, all reply tweets (with timestamps, reply-to links and user metadata), a "
            "structure.json conversation tree and an annotation.json label. Rumours were labelled by "
            "journalists as true, false or unverified.")
    rows = [[ev] + [int(summary.loc[ev, c]) for c in ("true", "false", "unverified", "nonrumour", "total")]
            for ev in summary.index]
    _table(doc, ["Event", "True", "False", "Unverified", "Non-rumour", "Total"], rows)
    _fig(doc, F / "fig1_dataset.png", "Figure 1. Threads per event and label.", 14)
    _p(doc, "Two tasks are defined. Task A (main): false vs true rumours "
            f"({int(summary.loc['TOTAL', 'false'])} vs {int(summary.loc['TOTAL', 'true'])} threads; "
            "unverified rumours are excluded because their ground truth is unknown). Task B (supporting): "
            f"rumour vs non-rumour over all {int(summary.loc['TOTAL', 'total'])} threads.")
    _p(doc, "PHEME records reply conversations, not retweets. Our cascades are therefore reply cascades: "
            "an edge means 'user B replied to user A's tweet'. Labels are strongly unbalanced across "
            "events (e.g. prince-toronto is almost entirely false, gurlitt entirely true), which is why "
            "pooled out-of-fold metrics are reported in addition to per-event results.",
       bold_lead="Note. ")

    # ---------------- 5. methodology ------------------------------------------
    doc.add_heading("5. Methodology", level=1)
    _fig(doc, F / "pipeline.png", "Figure 2. End-to-end pipeline.", 16)
    doc.add_heading("5.1 Network representation", level=2)
    _bullets(doc, [
        ("Cascade tree (tweet level, directed): ", "node = tweet, edge = parent tweet -> reply. "
         "Captures the geometry of the conversation."),
        ("User interaction graph (user level, directed, weighted): ", "node = user, edge = replier -> "
         "author of the replied-to tweet. Captures who talks to whom; used for centrality, reciprocity "
         "and Louvain community detection."),
        ("Node metadata: ", "timestamp, follower and friend counts, verified flag and account age at "
         "posting time. The tweet text is never used by the structural models."),
    ])
    doc.add_heading("5.2 Early-window truncation", level=2)
    _p(doc, "For every thread we keep only the tweets posted within t minutes of the source tweet, "
            f"for t in {{{', '.join(W[:-1])}}} minutes, plus the complete cascade. Replies whose parent "
            "is not yet visible are attached to the source, exactly as a real-time observer would see "
            "them. This simulates what an early-warning system would actually observe.")
    doc.add_heading("5.3 Structural features", level=2)
    rows = [[g, f, FEATURE_DESCRIPTIONS[f]] for g, fs in FEATURE_GROUPS.items() for f in fs]
    _table(doc, ["Group", "Feature", "Meaning"], rows, widths=[2.4, 4.6, 9.5], font=8.5)
    _p(doc, "Structural virality is the average shortest-path distance between all pairs of nodes in "
            "the cascade tree [8]: low for broadcast-like stars, high for long person-to-person chains. "
            "Communities are found with the Louvain modularity method [9].")
    doc.add_heading("5.4 Models and evaluation protocol", level=2)
    _bullets(doc, [
        ("Classifiers: ", "Logistic Regression (interpretable linear model), Random Forest and Histogram "
         "Gradient Boosting; class-balanced weighting; median imputation. The main model for each task is "
         f"chosen automatically as the one with the best mean AUC across windows (Task A: {main}; "
         f"Task B: {mains['B_rumour_vs_nonrumour']}); importance, ablation and the demo use it."),
        ("Leave-one-event-out CV: ", "train on 8 events, predict the held-out 9th, repeat for every "
         "event and pool the out-of-fold predictions. No event ever appears in both training and test."),
        ("Baselines: ", "(i) full-cascade SNA model (the 'full' window); (ii) content baseline - TF-IDF "
         "unigrams+bigrams of the source tweet with logistic regression, same CV protocol."),
        ("Metrics: ", "AUC-ROC (with 95% bootstrap CI), macro-F1 and accuracy at each window; "
         "permutation importance (drop in AUC on held-out events); time-to-flag."),
        ("Ablation: ", "topology-only features vs user-metadata-only features."),
    ])

    # ---------------- 6. results ----------------------------------------------
    doc.add_heading("6. Results", level=1)
    doc.add_heading("6.1 Early growth of true and false conversations", level=2)
    _fig(doc, F / "fig2_cascade_growth.png",
         "Figure 3. Mean cumulative replies per thread over the first three hours.")

    doc.add_heading("6.2 Detection performance over time (main result)", level=2)
    _fig(doc, F / "fig3_auc_vs_time_A.png",
         f"Figure 4. Task A AUC vs observation window. Shaded band: 95% bootstrap CI ({main}, the main model). "
         "Dashed: text baseline.")
    rows = []
    for w in W:
        rows.append([w if w == "full" else f"{w} min", f"{A['auc'][w]:.3f}",
                     f"[{A['lo'][w]:.3f}, {A['hi'][w]:.3f}]", f"{A['f1'][w]:.3f}",
                     _pct(A["retained"][w]), A["best_model"][w]])
    _table(doc, ["Window", f"AUC ({main})", "95% CI", "Macro-F1", "Signal retained*", "Best model"], rows)
    _p(doc, "*Signal retained = (AUC_t - 0.5) / (AUC_full - 0.5): the share of the full-cascade "
            "improvement over chance that is already available at time t.", size=9)
    beats = A["beats_text_from"]
    txt = (f"After 10 minutes the structural model already reaches AUC {A['auc']['10']:.2f}, and "
           f"after 30 minutes {A['auc']['30']:.2f} - {_pct(A['retained']['30'])} of the full-cascade "
           f"signal. Performance keeps rising with more observation time, reaching {A['auc']['full']:.2f} "
           "on the complete cascade. ")
    if beats:
        txt += (f"The structure-only model matches or exceeds the text baseline (AUC {A['text_auc']:.2f}) "
                f"from the {beats}{' min' if beats != 'full' else ''} window onwards, without reading a single word.")
    else:
        txt += (f"The text baseline (AUC {A['text_auc']:.2f}) remains stronger at every window, so structure "
                "is best seen as a complementary, early and hard-to-fake signal rather than a replacement for content.")
    _p(doc, txt)
    _fig(doc, F / "fig3_f1_vs_time_A.png", "Figure 5. Task A macro-F1 vs observation window.")
    _fig(doc, F / "fig3_auc_vs_time_B.png", "Figure 6. Task B (rumour vs non-rumour) AUC vs window.")
    _p(doc, f"Task B shows the same pattern on the larger sample: AUC {B['auc']['10']:.2f} at 10 minutes, "
            f"{B['auc']['30']:.2f} at 30 minutes and {B['auc']['full']:.2f} on the full cascade "
            f"(text baseline {B['text_auc']:.2f}).")

    doc.add_heading("6.3 Which signals carry the information? (ablation)", level=2)
    _fig(doc, F / "fig4_ablation_A.png", "Figure 7. Task A: topology-only vs user-metadata-only features.")
    topo30, meta30 = A["topo"]["30"], A["meta"]["30"]
    lead = "user metadata (who is talking)" if meta30 > topo30 else "pure topology (how the conversation is shaped)"
    _p(doc, f"At 30 minutes, topology-only features give AUC {topo30:.2f} and user-metadata-only features "
            f"give {meta30:.2f}; combined they reach {A['auc']['30']:.2f}. The stronger single source of "
            f"signal is {lead}, and combining both groups gives the best result, so the two views are "
            "complementary.")

    doc.add_heading("6.4 Feature importance across time", level=2)
    _fig(doc, F / "fig5_importance_A.png", "Figure 8. Permutation importance (drop in AUC on held-out events) per window.")
    _fig(doc, F / "fig6_distributions_A.png", "Figure 9. Distribution of the top features at 30 minutes.", 15)
    t10 = ", ".join(A["top"]["10"])
    tf = ", ".join(A["top"]["full"])
    _p(doc, f"The most informative features in the first 10 minutes are {t10}; on the full cascade they "
            f"are {tf}. Early decisions rely on fast-to-observe signals, while shape and community "
            "features gain weight as the conversation grows.")

    if "ttf" in K:
        doc.add_heading("6.5 How early are false rumours caught?", level=2)
        _fig(doc, F / "fig7_time_to_flag.png",
             "Figure 10. Cumulative share of false rumours caught (orange) and true rumours wrongly "
             "flagged (blue) by each window, threshold 0.5.")
        t, fa = K["ttf"], K.get("ttf_fa", {})
        _p(doc, f"{_pct(t['10'])} of false rumours are flagged within 10 minutes and {_pct(t['30'])} "
                f"within 30 minutes. The cost is false alarms: by 10 minutes {_pct(fa.get('10', 0))} of "
                f"true rumours have also crossed the threshold at least once. In practice the score is "
                "best used to rank claims for human fact-checkers, and the threshold can be raised to "
                "trade recall for fewer false alarms.")

    doc.add_heading("6.6 Generalisation to unseen events", level=2)
    _fig(doc, F / "fig8_per_event_A.png", "Figure 11. AUC on each held-out event (events with a single class omitted).")
    pe = per_event[(per_event.task == "A_false_vs_true") & (per_event.window == "30")].dropna(subset=["auc"])
    if len(pe):
        best, worst = pe.sort_values("auc").iloc[-1], pe.sort_values("auc").iloc[0]
        _p(doc, f"At 30 minutes, AUC ranges from {worst.auc:.2f} ({worst.event}) to {best.auc:.2f} "
                f"({best.event}). Variation across events is expected: each crisis has its own audience, "
                "pace and mix of true and false claims.")

    doc.add_heading("6.7 Summary of all models", level=2)
    m = metrics[(metrics.task == "A_false_vs_true") & (metrics.feature_set == "all")]
    piv = m.pivot_table(index="model", columns="window", values="auc").reindex(columns=W)
    rows = [[mdl] + [f"{piv.at[mdl, w]:.3f}" for w in W] for mdl in piv.index]
    rows.append(["Text baseline (TF-IDF)"] + [f"{A['text_auc']:.3f}"] * len(W))
    _table(doc, ["Model (Task A AUC)"] + [w if w == "full" else f"{w}m" for w in W], rows)

    # ---------------- 7. discussion ------------------------------------------
    doc.add_heading("7. Discussion", level=1)
    _bullets(doc, [
        ("Early signal is real. ", f"A large share of the predictive signal "
         f"({_pct(A['retained']['30'])} by 30 minutes) is available long before a cascade completes, "
         "supporting the project's central claim."),
        ("Structure is hard to fake. ", "Unlike wording, the timing, depth and audience of a "
         "conversation are produced by many independent users, which makes them harder to manipulate "
         "for a single author."),
        ("Interpretability. ", "Every decision can be explained with a handful of measurable network "
         "properties (see the 'why flagged' view in the demo), unlike an LLM text classifier."),
    ])
    doc.add_heading("Limitations", level=2)
    _bullets(doc, [
        "PHEME contains reply cascades, not retweet cascades; retweet diffusion may show different patterns.",
        "Many cascades are small; in the first 5-10 minutes several have only 0-3 replies, which limits "
        "centrality and community features in the earliest windows.",
        "Labels are unevenly distributed across events, so per-event AUC is undefined for single-class "
        "events and pooled metrics can be influenced by event mix.",
        "User metadata (followers, account age) was captured at collection time and is a proxy for the "
        "account state at posting time.",
        "The data is from 2014-2015 Twitter; platform behaviour has changed since.",
    ])

    doc.add_heading("8. Conclusion and Future Work", level=1)
    _p(doc, "We built an end-to-end pipeline that observes a partially-formed sharing cascade and "
            "estimates, from network structure alone, how likely it is to be misinformation. Evaluated on "
            "unseen events, the model achieves useful accuracy within the first 10-30 minutes and "
            "improves steadily as the cascade grows, quantifying the earliness/accuracy trade-off that "
            "existing full-cascade systems ignore.")
    _bullets(doc, [
        "Fuse the early structural score with a lightweight text model and test whether they are complementary.",
        "Use temporal graph neural networks over the snapshot sequence instead of hand-crafted features.",
        "Validate on retweet-based datasets (e.g. FakeNewsNet) and on non-English events.",
        "Deploy as a triage tool that ranks live claims for human fact-checkers.",
    ])

    doc.add_heading("References", level=1)
    for i, ref in enumerate(REFERENCES, 1):
        p = doc.add_paragraph(f"[{i}] {ref}")
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.size = Pt(9.5)

    doc.add_heading("Appendix A. Reproducing the results", level=1)
    _p(doc, f"Data source: {run_info.get('data_dir')}.  Threads: {run_info.get('n_threads')}.  "
            f"Tweets: {run_info.get('n_tweets')}.  Windows: {', '.join(W)}.  Runtime: "
            f"{run_info.get('runtime_min', '?')} min.")
    _bullets(doc, ["pip install -r requirements.txt",
                   "Place the extracted all-rnr-annotated-threads folder in data/raw/",
                   "python run_all.py   (data -> features -> models -> charts -> demo data -> this report)",
                   "streamlit run app/demo_app.py   (interactive demo)"])

    doc.save(out_file)
    return out_file
